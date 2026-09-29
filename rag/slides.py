"""Local, editable one-slide export. No LLM calls or invented numeric series."""
from io import BytesIO
import re
from pathlib import PurePosixPath

from rag.analysis import MODE_LABELS
from rag.evidence import REFUSAL

PPTX_MIME = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'


def _references(text):
    return sorted({int(n) for n in re.findall(r'\[S(\d+)\]', text)})


def can_export(result):
    answer = result.get('answer', '')
    refs = _references(answer)
    return bool(answer and answer != REFUSAL and 'INSUFFICIENT_EVIDENCE' not in answer
                and result.get('evidence_status') != 'insufficient' and refs
                and all(1 <= n <= len(result.get('sources', [])) for n in refs))


def _excerpts(answer):
    body = answer.split('\n\nИсточники:\n', 1)[0]
    # Keep complete cited paragraphs. Never cut a value, caveat or citation to fit.
    parts = [p.strip() for p in re.split(r'\n\s*\n|\n(?=[•\-]|\d+[.)] )', body) if p.strip()]
    chosen = [p for p in parts if _references(p) and len(p) <= 360 and len(p.splitlines()) <= 5][:3]
    return chosen, len(chosen) != len(parts)


def export_answer_pptx(result):
    if not can_export(result):
        raise ValueError('Для слайда нужен ответ с корректными ссылками на источники.')
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    navy, teal, gray = '16324F', '007F82', '526270'
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string('F7FAFC')

    def textbox(text, x, y, w, h, size=16, color=navy, bold=False, url=None):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = shape.text_frame
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        p = frame.paragraphs[0]
        run = p.add_run()
        run.text = text
        run.font.name, run.font.size, run.font.bold = 'Arial', Pt(size), bold
        run.font.color.rgb = RGBColor.from_string(color)
        if url:
            run.hyperlink.address = url
        return shape

    question = re.sub(r'\s+', ' ', result.get('question') or 'Аналитическая справка')
    title = question if len(question) <= 120 else question[:117].rsplit(' ', 1)[0] + '…'
    textbox(title, .55, .35, 12.2, .95, size=25, bold=True)
    mode = MODE_LABELS.get(result.get('mode'), MODE_LABELS['text'])
    created = str(result.get('generated_at') or 'не указана')
    textbox(f'{mode}  •  Сформировано: {created[:10]}', .55, 1.35, 12.2, .35, size=13, color=teal)
    excerpts, shortened = _excerpts(result['answer'])
    if not excerpts:
        excerpts = ['Подробный ответ превышает объём слайда. Полный текст и источники — в заметках.']
    table = slide.shapes.add_table(len(excerpts) + 1, 2, Inches(.55), Inches(1.95),
                                   Inches(12.2), Inches(3.65)).table
    table.columns[0].width, table.columns[1].width = Inches(9.7), Inches(2.5)
    table.rows[0].height = Inches(.45)
    table.cell(0, 0).text, table.cell(0, 1).text = 'Выводы и данные', 'Подтверждения'
    for i, excerpt in enumerate(excerpts, 1):
        table.rows[i].height = Inches(3.2 / len(excerpts))
        table.cell(i, 0).text = excerpt
        table.cell(i, 1).text = ', '.join(f'[S{n}]' for n in _references(excerpt)) or 'См. заметки'
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(navy if i == 0 else 'FFFFFF')
            cell.margin_left = cell.margin_right = Inches(.12)
            cell.margin_top = cell.margin_bottom = Inches(.07)
            cell.text_frame.word_wrap = True
            cell.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            for p in cell.text_frame.paragraphs:
                p.font.name, p.font.size = 'Arial', Pt(14 if i else 15)
                p.font.bold = i == 0
                p.font.color.rgb = RGBColor.from_string('FFFFFF' if i == 0 else navy)

    refs = _references(result['answer'])
    source_notes = []
    for i, n in enumerate(refs):
        meta = result['sources'][n - 1].metadata
        name = str(meta.get('title') or PurePosixPath(str(meta.get('source', 'Источник')).replace('\\', '/')).name)
        page = meta.get('page')
        location = f'стр. {page + 1}' if isinstance(page, int) else 'страница неизвестна'
        url = meta.get('source_url', '')
        source_notes.append(f'[S{n}] {name}, {location}; {meta.get("source_kind", "unknown")}\n{url}')
        if i < 8:
            label = name if len(name) <= 54 else name[:51] + '…'
            textbox(f'[S{n}] {label}, {location}', .55 + (i % 2) * 6.15,
                    6.0 + (i // 2) * .25, 6.0, .25, size=10, color=gray, url=url or None)
    notice = 'Выдержки; полный ответ и оговорки — в заметках.' if shortened else 'Полный ответ и источники — в заметках.'
    textbox(notice, .55, 5.7, 12.2, .25, size=11, color=gray)
    textbox('Автоматический ответ: ссылки проверены структурно; факты требуют проверки.',
            .55, 7.13, 12.2, .22, size=10, color=gray)
    slide.notes_slide.notes_text_frame.text = (
        f'Вопрос: {result.get("question", "")}\nРежим: {mode}\nСформировано: {created}\n'
        'Дата формирования не является датой актуальности показателей. Периоды указаны в ответе.\n\n'
        + result['answer'] + '\n\nИсточники слайда:\n' + '\n'.join(source_notes))
    output = BytesIO()
    deck.save(output)
    return output.getvalue()
