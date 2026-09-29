from io import BytesIO
from pathlib import Path
import sys
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / '.runtime_deps'), str(ROOT / '.test_deps')]
from langchain_core.documents import Document
from rag.slides import export_answer_pptx
from rag.evidence import REFUSAL


class SlidesTests(unittest.TestCase):
    def result(self):
        return {'question': 'Активы банка', 'mode': 'expert', 'generated_at': '2026-09-26T10:00:00+00:00',
                'answer': 'Активы — 42,75 млрд руб. [S1]\n\nДругая оценка — 43,10 млрд руб. [S2]',
                'evidence_status': 'cited_unverified', 'sources': [
                    Document(page_content='42,75', metadata={'source': 'report.pdf', 'page': 2,
                             'source_kind': 'official_report', 'source_url': 'https://bank.example/report.pdf'}),
                    Document(page_content='43,10', metadata={'source': 'review.pdf', 'page': 4})]}

    def test_single_editable_slide_preserves_values_sources_and_notes(self):
        from pptx import Presentation
        result = self.result()
        deck = Presentation(BytesIO(export_answer_pptx(result)))
        self.assertEqual(len(deck.slides), 1)
        slide = deck.slides[0]
        table = next(s.table for s in slide.shapes if s.has_table)
        text = '\n'.join(c.text for row in table.rows for c in row.cells)
        self.assertIn('42,75', text)
        self.assertIn('43,10', text)
        self.assertIn('[S2]', text)
        self.assertIn(result['answer'], slide.notes_slide.notes_text_frame.text)
        self.assertIn('https://bank.example/report.pdf', slide.notes_slide.notes_text_frame.text)
        self.assertIn('стр. 3', slide.notes_slide.notes_text_frame.text)
        self.assertIn('2026-09-26', slide.notes_slide.notes_text_frame.text)
        for shape in slide.shapes:
            self.assertLessEqual(shape.left + shape.width, deck.slide_width)
            self.assertLessEqual(shape.top + shape.height, deck.slide_height)

    def test_refusal_and_invalid_citations_cannot_be_exported_as_analysis(self):
        for answer in (REFUSAL, '42 [S9]', 'Без источников'):
            with self.assertRaises(ValueError):
                export_answer_pptx(dict(self.result(), answer=answer))

    def test_long_content_goes_to_notes_without_silently_cutting_figures(self):
        from pptx import Presentation
        result = self.result()
        result['answer'] = ('Длинное пояснение ' * 80) + '42,7531 млрд руб. [S1]'
        slide = Presentation(BytesIO(export_answer_pptx(result))).slides[0]
        self.assertIn(result['answer'], slide.notes_slide.notes_text_frame.text)
        texts = '\n'.join(s.text for s in slide.shapes if s.has_text_frame)
        self.assertIn('заметках', texts)


if __name__ == '__main__':
    unittest.main()
